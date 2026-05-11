import base64
import io
import json
from datasette import hookimpl
from datasette.utils.asgi import Response


@hookimpl
def register_output_renderer(datasette):
    renderers = []

    try:
        import openpyxl
        renderers.append({
            "extension": "xlsx",
            "render": render_excel,
            "can_render": lambda **kwargs: True,
        })
    except ImportError:
        pass

    try:
        import pyarrow
        renderers.append({
            "extension": "parquet",
            "render": render_parquet,
            "can_render": lambda **kwargs: True,
        })
    except ImportError:
        pass

    return renderers


def serialize_value(value):
    if value is None:
        return None
    if isinstance(value, (bytes, bytearray)):
        return base64.b64encode(value).decode("utf-8")
    if isinstance(value, (list, dict)):
        return json.dumps(value, default=str)
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value


async def render_excel(
    datasette,
    columns,
    rows,
    sql,
    query_name,
    database,
    table,
    request,
    view_name,
    data,
):
    import openpyxl
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = table or "Data"

    for col_idx, column in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.value = str(column)

    for row_idx, row in enumerate(rows, 2):
        for col_idx, value in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.value = serialize_value(value)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"{table or database}.xlsx"
    return Response(
        output.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
        },
    )


async def render_parquet(
    datasette,
    columns,
    rows,
    sql,
    query_name,
    database,
    table,
    request,
    view_name,
    data,
):
    import pyarrow as pa
    import pyarrow.parquet as pq

    column_names = [str(col) for col in columns]

    row_dicts = []
    for row in rows:
        row_dict = {}
        for col, val in zip(column_names, row):
            row_dict[col] = serialize_value(val)
        row_dicts.append(row_dict)

    table_data = pa.Table.from_pylist(row_dicts)

    output = io.BytesIO()
    pq.write_table(table_data, output)
    output.seek(0)

    filename = f"{table or database}.parquet"
    return Response(
        output.read(),
        content_type="application/vnd.apache.parquet",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
        },
    )
