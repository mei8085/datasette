import io
import pytest

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

try:
    import pyarrow as pa
    import pyarrow.parquet as pq
    HAS_PYARROW = True
except ImportError:
    HAS_PYARROW = False

from datasette.app import Datasette


@pytest.fixture
async def ds_client_with_data():
    ds = Datasette()
    db = ds.add_memory_database("test_db")
    await db.execute_write_script("""
        CREATE TABLE test_table (
            id INTEGER PRIMARY KEY,
            name TEXT,
            value REAL
        );
        INSERT INTO test_table (id, name, value) VALUES (1, 'Alice', 10.5);
        INSERT INTO test_table (id, name, value) VALUES (2, 'Bob', 20.5);
        INSERT INTO test_table (id, name, value) VALUES (3, 'Charlie', 30.5);
    """)
    await ds.invoke_startup()
    return ds.client


@pytest.mark.skipif(not HAS_OPENPYXL, reason="openpyxl not installed")
@pytest.mark.asyncio
async def test_excel_export(ds_client_with_data):
    response = await ds_client_with_data.get("/test_db/test_table.xlsx")
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert 'attachment; filename="test_table.xlsx"' in response.headers["content-disposition"]

    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    assert "test_table" in wb.sheetnames
    ws = wb["test_table"]
    
    headers = [cell.value for cell in ws[1]]
    assert headers == ["id", "name", "value"]
    
    data = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        data.append(row)
    
    assert data == [
        (1, "Alice", 10.5),
        (2, "Bob", 20.5),
        (3, "Charlie", 30.5),
    ]


@pytest.mark.skipif(not HAS_PYARROW, reason="pyarrow not installed")
@pytest.mark.asyncio
async def test_parquet_export(ds_client_with_data):
    response = await ds_client_with_data.get("/test_db/test_table.parquet")
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.apache.parquet"
    assert 'attachment; filename="test_table.parquet"' in response.headers["content-disposition"]

    table = pq.read_table(io.BytesIO(response.content))
    assert table.column_names == ["id", "name", "value"]
    assert table.num_rows == 3
    
    data = table.to_pylist()
    assert data == [
        {"id": 1, "name": "Alice", "value": 10.5},
        {"id": 2, "name": "Bob", "value": 20.5},
        {"id": 3, "name": "Charlie", "value": 30.5},
    ]


@pytest.mark.skipif(not HAS_OPENPYXL, reason="openpyxl not installed")
@pytest.mark.asyncio
async def test_excel_export_query(ds_client_with_data):
    response = await ds_client_with_data.get(
        "/test_db/-/query.xlsx?sql=SELECT+id,+name+FROM+test_table+WHERE+id+%3C+3"
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    ws = wb.active
    
    headers = [cell.value for cell in ws[1]]
    assert headers == ["id", "name"]
    
    data = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        data.append(row)
    
    assert data == [
        (1, "Alice"),
        (2, "Bob"),
    ]


@pytest.mark.skipif(not HAS_PYARROW, reason="pyarrow not installed")
@pytest.mark.asyncio
async def test_parquet_export_query(ds_client_with_data):
    response = await ds_client_with_data.get(
        "/test_db/-/query.parquet?sql=SELECT+id,+name+FROM+test_table+WHERE+id+%3C+3"
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.apache.parquet"

    table = pq.read_table(io.BytesIO(response.content))
    assert table.column_names == ["id", "name"]
    assert table.num_rows == 2
    
    data = table.to_pylist()
    assert data == [
        {"id": 1, "name": "Alice"},
        {"id": 2, "name": "Bob"},
    ]


@pytest.mark.skipif(not HAS_OPENPYXL, reason="openpyxl not installed")
@pytest.mark.asyncio
async def test_excel_renderer_in_renderers_list(ds_client_with_data):
    response = await ds_client_with_data.get("/test_db/test_table")
    assert response.status_code == 200
    assert ".xlsx" in response.text


@pytest.mark.skipif(not HAS_PYARROW, reason="pyarrow not installed")
@pytest.mark.asyncio
async def test_parquet_renderer_in_renderers_list(ds_client_with_data):
    response = await ds_client_with_data.get("/test_db/test_table")
    assert response.status_code == 200
    assert ".parquet" in response.text


@pytest.mark.skipif(not HAS_OPENPYXL, reason="openpyxl not installed")
@pytest.mark.asyncio
async def test_excel_export_complex_data():
    ds = Datasette()
    db = ds.add_memory_database("complex_db")
    await db.execute_write_script("""
        CREATE TABLE complex_table (
            id INTEGER PRIMARY KEY,
            json_data TEXT,
            null_field TEXT,
            timestamp TEXT
        );
        INSERT INTO complex_table (id, json_data, null_field, timestamp) 
        VALUES (1, '{"key": "value"}', NULL, '2023-01-01 12:00:00');
    """)
    await ds.invoke_startup()
    client = ds.client
    
    response = await client.get("/complex_db/complex_table.xlsx")
    assert response.status_code == 200

    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    ws = wb.active
    
    data = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        data.append(row)
    
    assert len(data) == 1
    assert data[0][0] == 1
    assert data[0][1] == '{"key": "value"}'
    assert data[0][2] is None


@pytest.mark.skipif(not HAS_PYARROW, reason="pyarrow not installed")
@pytest.mark.asyncio
async def test_parquet_export_complex_data():
    ds = Datasette()
    db = ds.add_memory_database("complex_db")
    await db.execute_write_script("""
        CREATE TABLE complex_table (
            id INTEGER PRIMARY KEY,
            json_data TEXT,
            null_field TEXT
        );
        INSERT INTO complex_table (id, json_data, null_field) 
        VALUES (1, '{"key": "value"}', NULL);
    """)
    await ds.invoke_startup()
    client = ds.client
    
    response = await client.get("/complex_db/complex_table.parquet")
    assert response.status_code == 200

    table = pq.read_table(io.BytesIO(response.content))
    data = table.to_pylist()
    
    assert len(data) == 1
    assert data[0]["id"] == 1
    assert data[0]["json_data"] == '{"key": "value"}'
    assert data[0]["null_field"] is None


@pytest.mark.skipif(not HAS_OPENPYXL, reason="openpyxl not installed")
@pytest.mark.asyncio
async def test_renderers_registered_in_core():
    ds = Datasette()
    await ds.invoke_startup()
    
    assert "xlsx" in ds.renderers
    assert "parquet" in ds.renderers or not HAS_PYARROW


@pytest.mark.skipif(not HAS_OPENPYXL, reason="openpyxl not installed")
@pytest.mark.asyncio
async def test_excel_export_with_existing_fixtures(ds_client):
    response = await ds_client.get("/fixtures/simple_primary_key.xlsx")
    assert response.status_code == 200
    
    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    ws = wb.active
    
    headers = [cell.value for cell in ws[1]]
    assert headers == ["id", "content"]
    
    data = []
    for row in ws.iter_rows(min_row=2, max_row=3, values_only=True):
        data.append(row)
    
    assert (1, "hello") in data
    assert (2, "world") in data
